"""Shared helpers for editable Chinese PPTX decks."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

NAVY = RGBColor(0x0F, 0x2C, 0x59)
TEAL = RGBColor(0x1F, 0x6F, 0x8B)
ORANGE = RGBColor(0xC4, 0x5C, 0x26)
RED = RGBColor(0xB0, 0x3A, 0x2E)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x33, 0x3A, 0x44)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
MUTED = RGBColor(0x5B, 0x64, 0x70)
FONT = "Microsoft YaHei"


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def set_run_font(run, size=16, bold=False, color=SLATE, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def add_blank(prs):
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def bar(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def textbox(slide, l, t, w, h, text, size=16, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    return tb


def bullets(slide, l, t, w, h, items, size=16, color=SLATE, bold_first=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = item
        set_run_font(run, size=size, bold=(bold_first and i == 0), color=color)
    return tb


def header(slide, title, subtitle=None):
    bar(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.92), NAVY)
    textbox(slide, Inches(0.4), Inches(0.18), Inches(12.4), Inches(0.45), title, size=24, bold=True, color=WHITE)
    if subtitle:
        textbox(slide, Inches(0.4), Inches(0.55), Inches(12.4), Inches(0.3), subtitle, size=12, color=RGBColor(0xC9, 0xD6, 0xE5))
    bar(slide, Inches(0), Inches(0.92), Inches(13.333), Inches(0.08), TEAL)


def footer(slide, page, total, label):
    textbox(slide, Inches(0.4), Inches(7.15), Inches(9), Inches(0.25), label, size=10, color=MUTED)
    textbox(slide, Inches(11.2), Inches(7.15), Inches(1.7), Inches(0.25), f"{page} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def add_table(slide, l, t, w, h, rows, col_w=None, header_fill=NAVY):
    n_rows = len(rows)
    n_cols = len(rows[0])
    table_shape = slide.shapes.add_table(n_rows, n_cols, l, t, w, h)
    table = table_shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            table.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if r == 0 or c > 0 else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(val)
            if r == 0:
                set_run_font(run, size=11, bold=True, color=WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            else:
                set_run_font(run, size=11, bold=False, color=SLATE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return table
