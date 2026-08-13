#!/usr/bin/env python3
"""Build required XLSX deliverables with live Excel formulas."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, PieChart
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule, CellIsRule
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker

OUT = Path("/workspace/office-deliverables")
OUT.mkdir(parents=True, exist_ok=True)

HDR = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
HDR_FILL = PatternFill("solid", fgColor="0F2C59")
TEAL_FILL = PatternFill("solid", fgColor="1F6F8B")
RED_FILL = PatternFill("solid", fgColor="F4C7C3")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
YEL_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = HDR
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = THIN


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def t11():
    wb = Workbook()
    raw = wb.active
    raw.title = "Raw"
    headers = ["Order_ID", "Date", "Region", "Rep", "Product", "Units", "Unit_Price", "Status"]
    raw.append(headers)
    style_header(raw, 1, 8)
    rows = [
        (1001, "2026-07-01", "North", "A. Chen", "Alpha", 10, 120, "Completed"),
        (1002, "7/1/2026", "north", "Alex Chen", "Beta", 5, 200, "Completed"),
        (1003, "2026/07/02", "SOUTH", "Maria Li", "Alpha", 8, 120, "Completed"),
        (1004, "02-Jul-2026", "South", "Maria Li", "Gamma", 3, 500, "Cancelled"),
        (1005, "2026-07-03", "East", "Sam Wu", "Alpha", -2, 120, "Refund"),
        (1006, "2026-07-03", "east", "Sam Wu", "Beta", 4, 200, "Completed"),
        (1007, "2026-07-04", "West", "Nina Ko", "Gamma", 2, 500, "Completed"),
        (1007, "2026-07-04", "West", "Nina Ko", "Gamma", 2, 500, "Completed"),
        (1008, "2026-07-05", "west", "Nina K.", "Alpha", 6, 120, "Completed"),
        (1009, "2026-07-05", "", "Jordan", "Beta", 7, 200, "Completed"),
        (1010, "2026-07-06", "North", "A Chen", "Alpha", 10, 120, "Completed"),
        (1011, "2026-07-06", "North", "Alex Chen", "Beta", 5, 200, "Pending"),
        (1012, "2026-07-07", "East", "Sam Wu", "Gamma", 1, 500, "Completed"),
        (1013, "2026-07-07", "South", "Maria Li", "Beta", 5, 200, "Completed"),
        (1014, "2026-07-08", "West", "Nina Ko", "Alpha", -3, 120, "Refund"),
    ]
    for r in rows:
        raw.append(list(r))
    autosize(raw, [12, 14, 12, 14, 12, 10, 12, 14])
    raw["J1"] = "NOTE"
    raw["J2"] = "15 rows including one exact duplicate of Order 1007. Do not delete similar-but-not-identical orders (1001 vs 1010)."

    # Clean: 14 unique rows mapping to Raw rows 2-8, 10-16 (skip Raw row 9 = dup 1007)
    clean = wb.create_sheet("Clean")
    ch = ["Order_ID", "Date_ISO", "Region_Raw", "Region", "Rep_Raw", "Rep", "Product", "Units", "Unit_Price", "Status", "Net_Revenue", "Raw_Row"]
    clean.append(ch)
    style_header(clean, 1, 12)
    src_rows = [2, 3, 4, 5, 6, 7, 8, 10, 11, 12, 13, 14, 15, 16]  # skip 9
    iso_map = {
        2: "2026-07-01", 3: "2026-07-01", 4: "2026-07-02", 5: "2026-07-02",
        6: "2026-07-03", 7: "2026-07-03", 8: "2026-07-04", 10: "2026-07-05",
        11: "2026-07-05", 12: "2026-07-06", 13: "2026-07-06", 14: "2026-07-07",
        15: "2026-07-07", 16: "2026-07-08",
    }
    for i, src in enumerate(src_rows, start=2):
        clean.cell(i, 1, f"=Raw!A{src}")
        clean.cell(i, 2, iso_map[src])
        clean.cell(i, 3, f"=Raw!C{src}")
        clean.cell(i, 4, f'=IF(TRIM(C{i})="","Unknown",UPPER(LEFT(TRIM(C{i}),1))&LOWER(MID(TRIM(C{i}),2,99)))')
        clean.cell(i, 5, f"=Raw!D{src}")
        clean.cell(i, 6, f'=IFS(OR(E{i}="A. Chen",E{i}="A Chen",E{i}="Alex Chen"),"Alex Chen",OR(E{i}="Nina K.",E{i}="Nina Ko"),"Nina Ko",TRUE,E{i})')
        clean.cell(i, 7, f"=Raw!E{src}")
        clean.cell(i, 8, f"=Raw!F{src}")
        clean.cell(i, 9, f"=Raw!G{src}")
        clean.cell(i, 10, f"=Raw!H{src}")
        clean.cell(i, 11, f'=IF(J{i}="Completed",H{i}*I{i},IF(J{i}="Refund",H{i}*I{i},0))')
        clean.cell(i, 12, src)
        for c in range(1, 13):
            clean.cell(i, c).border = THIN
            if c == 11:
                clean.cell(i, c).number_format = '#,##0'
    clean["N1"] = "Rules"
    clean["N2"] = "Drop exact duplicate of 1007 (Raw row 9). Keep 1001 and 1010 (different dates). Completed=Units*Price; Refund=Units*Price (units already negative); Cancelled/Pending=0."
    autosize(clean, [12, 14, 14, 12, 14, 14, 12, 10, 12, 14, 14, 12])

    sm = wb.create_sheet("Summary")
    sm.append(["Region", "Orders", "Net_Revenue"])
    style_header(sm, 1, 3)
    regions = ["East", "North", "South", "West", "Unknown"]
    for i, reg in enumerate(regions, start=2):
        sm.cell(i, 1, reg)
        sm.cell(i, 2, f'=COUNTIF(Clean!D:D,A{i})')
        sm.cell(i, 3, f'=SUMIF(Clean!D:D,A{i},Clean!K:K)')
        sm.cell(i, 3).number_format = '#,##0'
    sm["A7"] = "Total"
    sm["B7"] = "=SUM(B2:B6)"
    sm["C7"] = "=SUM(C2:C6)"
    sm["C7"].number_format = '#,##0'
    sm["A7"].font = Font(bold=True)
    sm["E1"] = "Check_Clean_Rows"
    sm["F1"] = "=COUNTA(Clean!A2:A15)"
    sm["E2"] = "Check_Net_vs_Clean"
    sm["F2"] = "=C7-SUM(Clean!K2:K15)"
    sm["E3"] = "Expected_Net"
    sm["F3"] = 9180
    sm["E4"] = "Match"
    sm["F4"] = '=IF(AND(F1=14,F2=0,C7=F3),"OK","CHECK")'
    chart = BarChart()
    chart.type = "col"
    chart.title = "Net Revenue by Region"
    chart.y_axis.title = "Net Revenue"
    chart.x_axis.title = "Region"
    chart.y_axis.scaling.min = 0
    data = Reference(sm, min_col=3, min_row=1, max_row=6)
    cats = Reference(sm, min_col=1, min_row=2, max_row=6)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.style = 10
    sm.add_chart(chart, "A9")
    autosize(sm, [16, 12, 16, 14, 22, 14])

    ex = wb.create_sheet("Exceptions")
    ex.append(["Issue", "Order_ID", "Field", "Raw_Value", "Action", "Formula_Flag"])
    style_header(ex, 1, 6)
    ex["A2"] = "Missing Region"
    ex["B2"] = "=Clean!A10"
    ex["C2"] = "Region"
    ex["D2"] = "(blank)"
    ex["E2"] = "Coded Unknown; sales must complete"
    ex["F2"] = '=COUNTIF(Clean!D:D,"Unknown")'
    ex["A3"] = "Duplicate dropped"
    ex["B3"] = 1007
    ex["C3"] = "entire row"
    ex["D3"] = "identical second row (Raw row 9)"
    ex["E3"] = "Keep one"
    ex["F3"] = '=COUNTIF(Raw!A:A,1007)'
    ex["A5"] = "Refund rows are not exceptions; negative units follow the net-revenue rule."
    autosize(ex, [22, 12, 14, 36, 36, 18])
    wb.save(OUT / "T11_Sales_Orders_Clean_Summary.xlsx")


def t12():
    wb = Workbook()
    d = wb.active
    d.title = "Monthly_Detail"
    d.append(["Department", "Jan_B", "Jan_A", "Feb_B", "Feb_A", "Mar_B", "Mar_A", "Jan_Var", "Feb_Var", "Mar_Var"])
    style_header(d, 1, 10)
    data = [
        ("Marketing", 100, 105, 110, 98, 120, 150),
        ("R&D", 200, 190, 210, 230, 220, 250),
        ("G&A", 80, 78, 80, 82, 85, 80),
        ("Sales", 150, 145, 160, 175, 170, 185),
    ]
    for i, row in enumerate(data, start=2):
        d.cell(i, 1, row[0])
        for j, v in enumerate(row[1:], start=2):
            d.cell(i, j, v)
        d.cell(i, 8, f"=C{i}-B{i}")
        d.cell(i, 9, f"=E{i}-D{i}")
        d.cell(i, 10, f"=G{i}-F{i}")
    d["A6"] = "TOTAL"
    for col in range(2, 11):
        letter = get_column_letter(col)
        d.cell(6, col, f"={letter}2+{letter}3+{letter}4+{letter}5")
    d["A8"] = "Variance = Actual-Budget. Positive = Unfavorable for expenses. Unit: $000s."
    autosize(d, [14] + [10] * 9)

    s = wb.create_sheet("Dept_Summary")
    s.append(["Department", "Q_Budget", "Q_Actual", "Q_Var", "Q_Var_Pct", "Flag", "CumVar_Jan", "CumVar_Feb", "CumVar_Mar"])
    style_header(s, 1, 9)
    for i in range(2, 6):
        s.cell(i, 1, f"=Monthly_Detail!A{i}")
        s.cell(i, 2, f"=Monthly_Detail!B{i}+Monthly_Detail!D{i}+Monthly_Detail!F{i}")
        s.cell(i, 3, f"=Monthly_Detail!C{i}+Monthly_Detail!E{i}+Monthly_Detail!G{i}")
        s.cell(i, 4, f"=C{i}-B{i}")
        s.cell(i, 5, f"=D{i}/B{i}")
        s.cell(i, 5).number_format = "0.00%"
        s.cell(i, 6, f'=IF(D{i}<0,"Favorable-Green",IF(OR(D{i}>20,D{i}/B{i}>0.05),"Unfavorable-Red","Watch"))')
        s.cell(i, 7, f"=Monthly_Detail!H{i}")
        s.cell(i, 8, f"=G{i}+Monthly_Detail!I{i}")
        s.cell(i, 9, f"=H{i}+Monthly_Detail!J{i}")
    s["A6"] = "TOTAL"
    s["B6"] = "=SUM(B2:B5)"
    s["C6"] = "=SUM(C2:C5)"
    s["D6"] = "=C6-B6"
    s["E6"] = "=D6/B6"
    s["E6"].number_format = "0.00%"
    s.conditional_formatting.add("F2:F5", FormulaRule(formula=['ISNUMBER(SEARCH("Red",F2))'], fill=RED_FILL))
    s.conditional_formatting.add("F2:F5", FormulaRule(formula=['ISNUMBER(SEARCH("Green",F2))'], fill=GREEN_FILL))
    chart = BarChart()
    chart.type = "col"
    chart.title = "Quarter variance by department (Actual-Budget)"
    chart.y_axis.title = "$000s"
    data_ref = Reference(s, min_col=4, min_row=1, max_row=5)
    cats = Reference(s, min_col=1, min_row=2, max_row=5)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    s.add_chart(chart, "A9")
    line = LineChart()
    line.title = "Cumulative variance trend"
    line.y_axis.title = "$000s"
    line_data = Reference(s, min_col=7, min_row=1, max_col=9, max_row=5)
    line.add_data(line_data, from_rows=True, titles_from_data=False)
    line.set_categories(Reference(s, min_col=7, min_row=1, max_col=9))
    # Simpler line: departments as series across Jan/Feb/Mar
    line2 = LineChart()
    line2.title = "Cumulative variance (Jan / Feb / Mar)"
    cats2 = Reference(s, min_col=7, min_row=1, max_col=9, max_row=1)
    # Use a helper block for chart categories
    s["K1"] = "Jan"
    s["L1"] = "Feb"
    s["M1"] = "Mar"
    for i in range(2, 6):
        s.cell(i, 11, f"=G{i}")
        s.cell(i, 12, f"=H{i}")
        s.cell(i, 13, f"=I{i}")
    s["J2"] = "=A2"
    s["J3"] = "=A3"
    s["J4"] = "=A4"
    s["J5"] = "=A5"
    line2.y_axis.scaling  # keep 0 visible via data
    dref = Reference(s, min_col=11, min_row=1, max_col=13, max_row=5)
    line2.add_data(dref, from_rows=True, titles_from_data=True)
    line2.set_categories(Reference(s, min_col=11, min_row=1, max_col=13, max_row=1))
    # actually from_rows with header in K1:M1 and data K2:M5 needs series names from J
    s.add_chart(line2, "I9")
    c = wb.create_sheet("Conclusions")
    c["A1"] = "Three conclusions for management"
    c["A1"].font = Font(bold=True, size=14, color="0F2C59")
    c["A3"] = "1. R&D has the largest overspend: quarterly unfavorable 40 ($000s, 6.35%), with consecutive unfavorable months in Feb (+20) and Mar (+30). Do not be misled by January's favorable -10."
    c["A4"] = "2. Marketing is red because of March: Jan-Feb still favorable -7, then March +30 drives the quarter to +23 (6.97%). The issue is March spend, not even overspend."
    c["A5"] = "3. Sales is over the line (+25, 5.21%) for two consecutive months; G&A is favorable -5 (-2.04%) green. Company total +83 / 1685 = 4.93% would hide three red departments."
    c["A7"] = "Threshold: red if quarterly unfavorable >20 or variance rate >5% of department quarterly budget. Green if favorable."
    c["A8"] = "Formulas live on Monthly_Detail and Dept_Summary. Amounts in $000s."
    for r in range(3, 9):
        c.row_dimensions[r].height = 36
        c.cell(r, 1).alignment = Alignment(wrap_text=True)
    c.column_dimensions["A"].width = 120
    autosize(s, [14, 12, 12, 12, 12, 22, 14, 14, 14])
    wb.save(OUT / "T12_Q1_Department_Budget_Variance.xlsx")


def t13():
    wb = Workbook()
    inp = wb.active
    inp.title = "Input"
    inp.append(["Cohort", "Signups_M0", "M1_Active", "M2_Active", "M3_Active", "M4_Active"])
    style_header(inp, 1, 6)
    inp.append(["2026-Jan", 1000, 720, 610, 550, 500])
    inp.append(["2026-Feb", 1200, 840, 690, 600, None])
    inp.append(["2026-Mar", 900, 630, 520, None, None])
    inp.append(["2026-Apr", 1100, 770, None, None, None])
    inp["A7"] = "Unobserved months MUST remain blank, never 0."

    rates = wb.create_sheet("Rates")
    rates.append(["Cohort", "N", "M0", "M1", "M2", "M3", "M4"])
    style_header(rates, 1, 7)
    for i in range(2, 6):
        rates.cell(i, 1, f"=Input!A{i}")
        rates.cell(i, 2, f"=Input!B{i}")
        rates.cell(i, 3, 1)
        for col, src in [(4, "C"), (5, "D"), (6, "E"), (7, "F")]:
            rates.cell(i, col, f'=IF(Input!{src}{i}="","",Input!{src}{i}/$B{i})')
            rates.cell(i, col).number_format = "0.0%"
        rates.cell(i, 3).number_format = "0.0%"
    rates.conditional_formatting.add(
        "D2:G5",
        ColorScaleRule(start_type="num", start_value=0.5, start_color="F8696B",
                       mid_type="num", mid_value=0.65, mid_color="FFEB84",
                       end_type="num", end_value=0.75, end_color="63BE7B"),
    )

    w = wb.create_sheet("Weighted")
    w.append(["Month", "Active_Sum", "N_Sum", "Weighted_Retention", "Cohorts_Included"])
    style_header(w, 1, 5)
    w["A2"] = "M0"
    w["B2"] = "=SUM(Input!B2:B5)"
    w["C2"] = "=SUM(Input!B2:B5)"
    w["D2"] = "=B2/C2"
    w["E2"] = "Jan-Apr"
    w["A3"] = "M1"
    w["B3"] = "=SUM(Input!C2:C5)"
    w["C3"] = "=SUM(Input!B2:B5)"
    w["D3"] = "=B3/C3"
    w["E3"] = "Jan-Apr"
    w["A4"] = "M2"
    w["B4"] = "=SUM(Input!D2:D4)"
    w["C4"] = "=SUM(Input!B2:B4)"
    w["D4"] = "=B4/C4"
    w["E4"] = "Jan-Mar (Apr blank, excluded)"
    w["A5"] = "M3"
    w["B5"] = "=SUM(Input!E2:E3)"
    w["C5"] = "=SUM(Input!B2:B3)"
    w["D5"] = "=B5/C5"
    w["E5"] = "Jan-Feb"
    w["A6"] = "M4"
    w["B6"] = "=Input!F2"
    w["C6"] = "=Input!B2"
    w["D6"] = "=B6/C6"
    w["E6"] = "Jan only"
    for r in range(2, 7):
        w.cell(r, 4).number_format = "0.00%"
    w["A8"] = "Do NOT simple-average percentages. Weighted = sum of actives / sum of signups for available cohorts only."
    chart = LineChart()
    chart.title = "Weighted retention curve"
    chart.y_axis.title = "Retention"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.add_data(Reference(w, min_col=4, min_row=1, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(w, min_col=1, min_row=2, max_row=6))
    w.add_chart(chart, "A10")

    c = wb.create_sheet("Conclusion")
    c["A1"] = "结论（≤250字）"
    c["A1"].font = Font(bold=True, size=14)
    c["A3"] = ("M1 留存：Jan 72%，其后三个 cohort 均为 70%，没有出现引导上线后 M1 改善。"
               "2 月 15 日新手引导对 Feb 只覆盖半个月，Mar/Apr 应更充分暴露，但 M1 仍为 70%，低于 Jan。"
               "M2：Jan 61.0% vs Feb 57.5%、Mar 57.8%，后期队列更差。"
               "3 月 10 日折扣 30%→15% 与 Mar 部分、Apr 全部暴露同期，但无实验、无用户级控制，"
               "不能把留存走弱归因于引导无效或折扣下调。"
               "目前只能描述：新队列 M1/M2 弱于 Jan，趋势不支持“引导已改善留存”的管理叙事。")
    c["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    c.row_dimensions[3].height = 90
    c.column_dimensions["A"].width = 110
    autosize(inp, [14, 14, 14, 14, 14, 14])
    autosize(rates, [14, 10, 10, 10, 10, 10, 10])
    autosize(w, [12, 14, 12, 22, 36])
    wb.save(OUT / "T13_Cohort_Retention.xlsx")


def t14():
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    headers = ["SKU", "On_Hand", "Avg_Daily_Demand", "Lead_Time_Days", "Safety_Stock", "Open_PO", "Case_Pack",
               "ROP", "Available", "Trigger", "Target", "RawQty", "OrderQty"]
    ws.append(headers)
    style_header(ws, 1, 13)
    data = [
        ("A", 120, 8, 10, 40, 0, 20),
        ("B", 500, 12, 7, 60, 100, 50),
        ("C", 40, 5, 14, 25, 0, 10),
        ("D", 200, 20, 5, 50, 0, 25),
        ("E", 75, 4, 21, 30, 50, 10),
        ("F", 0, 3, 30, 20, 30, 25),
        ("G", 90, 6, 12, 36, 0, 12),
        ("H", 1000, 25, 14, 100, 500, 100),
    ]
    for i, row in enumerate(data, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(i, c, v)
        ws.cell(i, 8, f"=C{i}*D{i}+E{i}")
        ws.cell(i, 9, f"=B{i}+F{i}")
        ws.cell(i, 10, f'=IF(I{i}<=H{i},"YES","NO")')
        ws.cell(i, 11, f"=C{i}*(D{i}+14)+E{i}")
        ws.cell(i, 12, f"=MAX(0,K{i}-B{i}-F{i})")
        ws.cell(i, 13, f'=IF(J{i}="NO",0,IF(L{i}=0,0,CEILING(L{i},G{i})))')
    ws.conditional_formatting.add("J2:J9", FormulaRule(formula=['J2="YES"'], fill=RED_FILL))
    ws["A12"] = "ROP=AvgDaily*LeadTime+Safety. Trigger if OnHand+OpenPO <= ROP. Target=Demand*(LT+14)+Safety. Order=CEILING to case pack only if triggered."

    po = wb.create_sheet("Purchase_List")
    po.append(["SKU", "OrderQty", "Cases", "Priority_Note"])
    style_header(po, 1, 4)
    po["A2"] = '=IF(Inventory!J2="YES",Inventory!A2,"")'
    po["B2"] = '=IF(Inventory!J2="YES",Inventory!M2,"")'
    # Filter-style list: copy all triggered via formulas
    for i, src in enumerate(range(2, 10), start=2):
        po.cell(i, 1, f'=IF(Inventory!J{src}="YES",Inventory!A{src},"")')
        po.cell(i, 2, f'=IF(Inventory!J{src}="YES",Inventory!M{src},"")')
        po.cell(i, 3, f'=IF(B{i}="","",B{i}/INDEX(Inventory!G:G,MATCH(A{i},Inventory!A:A,0)))')
        po.cell(i, 4, f'=IF(A{i}="F","Highest stockout risk (OH=0)",IF(A{i}="","","Triggered"))')
    po["A11"] = "SKU count triggered"
    po["B11"] = '=COUNTIF(Inventory!J2:J9,"YES")'
    po["A12"] = "Units ordered"
    po["B12"] = "=SUM(Inventory!M2:M9)"

    ex = wb.create_sheet("Exception_Checks")
    ex.append(["Check", "SKU", "Result"])
    style_header(ex, 1, 3)
    notes = [
        ("On Hand = 0", "F", "Avail 30 <= ROP 110; Open PO still far below ROP"),
        ("Avail equals ROP", "A", "120=120 triggers because rule is <="),
        ("Below target but not ROP", "D", "200>150 ROP so do NOT order despite Target 430"),
        ("Below target but not ROP", "E", "125>114 ROP so do NOT order"),
        ("Far above ROP", "B/H", "Do not order"),
    ]
    for i, n in enumerate(notes, start=2):
        ex.cell(i, 1, n[0]); ex.cell(i, 2, n[1]); ex.cell(i, 3, n[2])
    autosize(ws, [8, 12, 18, 16, 14, 12, 12, 10, 12, 12, 12, 12, 12])
    autosize(po, [10, 12, 10, 40])
    autosize(ex, [28, 10, 70])
    wb.save(OUT / "T14_Reorder_Calculator.xlsx")


def t15():
    wb = Workbook()
    a = wb.active
    a.title = "Assumptions"
    a.append(["Item", "M1", "M2", "M3", "M4", "M5", "M6"])
    style_header(a, 1, 7)
    a.append(["Base_Revenue", 1.2, 1.3, 1.5, 1.4, 1.6, 1.8])
    a.append(["Downside_Factor", 0.85, 0.85, 0.85, 0.85, 0.85, 0.85])
    a["A4"] = "VC_Rate"
    a["B4"] = 0.32
    a.merge_cells("B4:G4")
    a.append(["Fixed", 0.75, 0.75, 0.75, 0.75, 0.75, 0.75])
    a.append(["Capex", 0.20, 0, 0.50, 0, 0.30, 0])
    a.append(["Financing", 0, 0.50, 0, 0, 0, 0])
    a["A8"] = "Opening_Cash_M1"
    a["B8"] = 2.0
    a["A9"] = "Min_Cash"
    a["B9"] = 0.8
    a["A10"] = "Downside_Revenue_M1"
    a["B10"] = "=B2*B3"
    a["C10"] = "=C2*C3"
    a["D10"] = "=D2*D3"
    a["E10"] = "=E2*E3"
    a["F10"] = "=F2*F3"
    a["G10"] = "=G2*G3"
    for col in range(2, 8):
        for row in range(2, 11):
            if a.cell(row, col).value is not None:
                a.cell(row, col).number_format = "0.000"

    def cash_sheet(name, rev_row_on_assumptions, opening_cell="Assumptions!B8"):
        ws = wb.create_sheet(name)
        ws.append(["Line", "M1", "M2", "M3", "M4", "M5", "M6"])
        style_header(ws, 1, 7)
        ws["A2"] = "Opening"
        ws["B2"] = f"={opening_cell}"
        ws["A3"] = "Revenue"
        for i, col in enumerate("BCDEFG"):
            src_col = get_column_letter(i + 2)
            ws.cell(3, i + 2, f"=Assumptions!{src_col}{rev_row_on_assumptions}")
        ws["A4"] = "VC"
        for i in range(2, 8):
            cl = get_column_letter(i)
            ws.cell(4, i, f"={cl}3*Assumptions!$B$4")
        ws["A5"] = "Fixed"
        for i, col in enumerate("BCDEFG"):
            ws.cell(5, i + 2, f"=Assumptions!{col}5")
        ws["A6"] = "Capex"
        for i, col in enumerate("BCDEFG"):
            ws.cell(6, i + 2, f"=Assumptions!{col}6")
        ws["A7"] = "Financing"
        for i, col in enumerate("BCDEFG"):
            ws.cell(7, i + 2, f"=Assumptions!{col}7")
        ws["A8"] = "Ending"
        for i in range(2, 8):
            cl = get_column_letter(i)
            ws.cell(8, i, f"={cl}2+{cl}3-{cl}4-{cl}5-{cl}6+{cl}7")
        for i in range(3, 8):
            prev = get_column_letter(i - 1)
            ws.cell(2, i, f"={prev}8")
        ws["A9"] = "vs_Min"
        for i in range(2, 8):
            cl = get_column_letter(i)
            ws.cell(9, i, f'=IF({cl}8>=Assumptions!$B$9,"OK","BREACH")')
        for r in range(2, 9):
            for c in range(2, 8):
                ws.cell(r, c).number_format = "0.000"
        return ws

    base = cash_sheet("Base", 2)
    down = cash_sheet("Downside", 10)

    dash = wb.create_sheet("Dashboard")
    dash["A1"] = "Six-month cash dashboard"
    dash["A1"].font = Font(bold=True, size=14, color="0F2C59")
    dash.append([])
    dash["A3"] = "Month"
    for i, m in enumerate(["M1", "M2", "M3", "M4", "M5", "M6"], start=2):
        dash.cell(3, i, m)
    dash["A4"] = "Base_Ending"
    dash["A5"] = "Downside_Ending"
    dash["A6"] = "Min_Cash"
    for i in range(2, 8):
        cl = get_column_letter(i)
        dash.cell(4, i, f"=Base!{cl}8")
        dash.cell(5, i, f"=Downside!{cl}8")
        dash.cell(6, i, f"=Assumptions!$B$9")
        dash.cell(4, i).number_format = "0.000"
        dash.cell(5, i).number_format = "0.000"
        dash.cell(6, i).number_format = "0.000"
    dash["A8"] = "Base_M6"
    dash["B8"] = "=Base!G8"
    dash["A9"] = "Downside_M6"
    dash["B9"] = "=Downside!G8"
    dash["A10"] = "Tightest_Base"
    dash["B10"] = "=MIN(B4:G4)"
    dash["A11"] = "Tightest_Downside"
    dash["B11"] = "=MIN(B5:G5)"
    dash["A12"] = "Any_breach"
    dash["B12"] = '=IF(OR(COUNTIF(Base!B9:G9,"BREACH")>0,COUNTIF(Downside!B9:G9,"BREACH")>0),"YES","NO")'
    dash["A14"] = "Both scenarios stay above 0.800. Tightest downside months are M1 and M5. M2 financing 0.500 is a key injection in both cases."
    chart = LineChart()
    chart.title = "Month-end cash vs 0.800 floor"
    chart.y_axis.title = "CNY millions"
    chart.y_axis.scaling.min = 0
    chart.add_data(Reference(dash, min_col=1, min_row=4, max_col=7, max_row=6), from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(dash, min_col=2, min_row=3, max_col=7, max_row=3))
    dash.add_chart(chart, "A16")
    autosize(a, [22, 10, 10, 10, 10, 10, 10])
    wb.save(OUT / "T15_Six_Month_Cash_Model.xlsx")


def t34():
    wb = Workbook()
    a = wb.active
    a.title = "Assumptions"
    a["A1"] = "Item"
    a["B1"] = "Value"
    a["C1"] = "Excel / note"
    style_header(a, 1, 3)
    rows = [
        ("Start", "2027-01-01", "Month-end payments"),
        ("N", 24, "periods"),
        ("Payment", 100000, "in arrears"),
        ("Annual_rate", 0.06, "effective annual"),
        ("Monthly_r", "=(1+B4)^(1/12)-1", "not 6%/12"),
        ("Initial_liability", "=B3*(1-(1+B5)^-B2)/B5", "ordinary annuity PV"),
        ("ROU_initial", "=B6", "equals initial liability"),
        ("Monthly_dep", "=B7/B2", "straight line 24 months"),
        ("VAT_IDC_prepaid_incentive_restore", "None", "per brief"),
    ]
    labels = ["Start", "N", "Payment", "Annual_rate", "Monthly_r", "Initial_liability", "ROU_initial", "Monthly_dep", "Exclusions"]
    a["A2"] = "Start"; a["B2"] = "2027-01-01"; a["C2"] = "Month-end payments"
    a["A3"] = "N"; a["B3"] = 24; a["C3"] = "periods"
    a["A4"] = "Payment"; a["B4"] = 100000; a["C4"] = "in arrears"
    a["A5"] = "Annual_rate"; a["B5"] = 0.06; a["B5"].number_format = "0.00%"; a["C5"] = "effective annual"
    a["A6"] = "Monthly_r"; a["B6"] = "=(1+B5)^(1/12)-1"; a["B6"].number_format = "0.0000000000"; a["C6"] = "=(1+6%)^(1/12)-1"
    a["A7"] = "Initial_liability"; a["B7"] = "=B4*(1-(1+B6)^-B3)/B6"; a["B7"].number_format = '#,##0.00'; a["C7"] = "PV ordinary annuity"
    a["A8"] = "ROU_initial"; a["B8"] = "=B7"; a["B8"].number_format = '#,##0.00'; a["C8"] = "equals liability"
    a["A9"] = "Monthly_dep"; a["B9"] = "=B8/B3"; a["B9"].number_format = '#,##0.00'; a["C9"] = "straight line"
    a["A10"] = "Display_round"; a["B10"] = "Yuan; keep precision in formulas"; a["C10"] = "do not chain rounded values"

    sch = wb.create_sheet("Schedule")
    sch.append(["Month", "Opening_Liability", "Interest", "Payment", "Principal", "Closing_Liability", "ROU_Depreciation", "Year"])
    style_header(sch, 1, 8)
    sch["A2"] = 1
    sch["B2"] = "=Assumptions!B7"
    sch["C2"] = "=B2*Assumptions!$B$6"
    sch["D2"] = "=Assumptions!$B$4"
    sch["E2"] = "=D2-C2"
    sch["F2"] = "=B2+C2-D2"
    sch["G2"] = "=Assumptions!$B$9"
    sch["H2"] = 2027
    for m in range(3, 26):
        prev = m - 1
        sch.cell(m, 1, m - 1)
        sch.cell(m, 2, f"=F{prev}")
        sch.cell(m, 3, f"=B{m}*Assumptions!$B$6")
        sch.cell(m, 4, "=Assumptions!$B$4")
        sch.cell(m, 5, f"=D{m}-C{m}")
        sch.cell(m, 6, f"=B{m}+C{m}-D{m}")
        sch.cell(m, 7, "=Assumptions!$B$9")
        sch.cell(m, 8, 2027 if m <= 13 else 2028)
    for r in range(2, 26):
        for c in range(2, 8):
            sch.cell(r, c).number_format = '#,##0.00'
    sch["A28"] = "Month 24 closing should be ~0. Internal formulas keep full precision."

    an = wb.create_sheet("Annual_Summary")
    an.append(["Year", "Interest", "Cash_Payments", "Depreciation", "Closing_Liability"])
    style_header(an, 1, 5)
    an["A2"] = 2027
    an["B2"] = "=SUMIF(Schedule!H:H,2027,Schedule!C:C)"
    an["C2"] = "=SUMIF(Schedule!H:H,2027,Schedule!D:D)"
    an["D2"] = "=SUMIF(Schedule!H:H,2027,Schedule!G:G)"
    an["E2"] = "=INDEX(Schedule!F:F,MATCH(12,Schedule!A:A,0))"
    an["A3"] = 2028
    an["B3"] = "=SUMIF(Schedule!H:H,2028,Schedule!C:C)"
    an["C3"] = "=SUMIF(Schedule!H:H,2028,Schedule!D:D)"
    an["D3"] = "=SUMIF(Schedule!H:H,2028,Schedule!G:G)"
    an["E3"] = "=INDEX(Schedule!F:F,MATCH(24,Schedule!A:A,0))"
    for r in range(2, 4):
        for c in range(2, 6):
            an.cell(r, c).number_format = '#,##0.00'
    an["A5"] = "Check_interest_plus_principal_vs_payments_24m"
    an["B5"] = "=SUM(Schedule!C2:C25)+SUM(Schedule!E2:E25)"
    an["C5"] = "=SUM(Schedule!D2:D25)"
    an["D5"] = "=B5-C5"
    autosize(a, [22, 28, 40])
    autosize(sch, [10, 20, 14, 12, 14, 20, 18, 10])
    autosize(an, [18, 16, 16, 16, 20])
    wb.save(OUT / "T34_IFRS16_Lease_Amortization.xlsx")


def t35():
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp.append(["Item", "Vendor_A", "Vendor_B", "Vendor_C", "Timing"])
    style_header(inp, 1, 5)
    rows = [
        ("Implementation_t0", 220, 150, 300, "t=0"),
        ("Training_t0", 60, 90, 40, "t=0"),
        ("Y1_License", 380, 300, 450, "year-end"),
        ("Y2_License", 380, 315, 450, "year-end"),
        ("Y3_License", 380, 331, 450, "year-end"),
        ("Annual_Support", 40, 65, 25, "each year-end"),
        ("Annual_Downtime", 25, 80, 10, "each year-end"),
        ("Y1_Migration", 0, 50, 0, "end of Y1"),
        ("Y3_Residual_inflow", 0, 0, 120, "end of Y3 inflow"),
        ("Discount_rate", 0.08, 0.08, 0.08, ""),
    ]
    for r in rows:
        inp.append(list(r))
    inp["B11"].number_format = "0.00%"
    inp["C11"].number_format = "0.00%"
    inp["D11"].number_format = "0.00%"

    cf = wb.create_sheet("CashFlows")
    cf.append(["Line", "A", "B", "C"])
    style_header(cf, 1, 4)
    cf["A2"] = "t0"
    cf["B2"] = "=Inputs!B2+Inputs!B3"
    cf["C2"] = "=Inputs!C2+Inputs!C3"
    cf["D2"] = "=Inputs!D2+Inputs!D3"
    cf["A3"] = "Y1"
    cf["B3"] = "=Inputs!B4+Inputs!B7+Inputs!B8+Inputs!B9"
    cf["C3"] = "=Inputs!C4+Inputs!C7+Inputs!C8+Inputs!C9"
    cf["D3"] = "=Inputs!D4+Inputs!D7+Inputs!D8+Inputs!D9"
    cf["A4"] = "Y2"
    cf["B4"] = "=Inputs!B5+Inputs!B7+Inputs!B8"
    cf["C4"] = "=Inputs!C5+Inputs!C7+Inputs!C8"
    cf["D4"] = "=Inputs!D5+Inputs!D7+Inputs!D8"
    cf["A5"] = "Y3_net"
    cf["B5"] = "=Inputs!B6+Inputs!B7+Inputs!B8-Inputs!B10"
    cf["C5"] = "=Inputs!C6+Inputs!C7+Inputs!C8-Inputs!C10"
    cf["D5"] = "=Inputs!D6+Inputs!D7+Inputs!D8-Inputs!D10"
    cf["A6"] = "Nominal_3yr"
    cf["B6"] = "=B2+B3+B4+B5"
    cf["C6"] = "=C2+C3+C4+C5"
    cf["D6"] = "=D2+D3+D4+D5"
    cf["A7"] = "NPV_8pct"
    cf["B7"] = "=B2+B3/(1+Inputs!B11)+B4/(1+Inputs!B11)^2+B5/(1+Inputs!B11)^3"
    cf["C7"] = "=C2+C3/(1+Inputs!C11)+C4/(1+Inputs!C11)^2+C5/(1+Inputs!C11)^3"
    cf["D7"] = "=D2+D3/(1+Inputs!D11)+D4/(1+Inputs!D11)^2+D5/(1+Inputs!D11)^3"
    for r in range(2, 8):
        for c in range(2, 5):
            cf.cell(r, c).number_format = '#,##0.00'
    cf["A9"] = "Lowest_NPV"
    cf["B9"] = '=INDEX(A7:D7,1,MATCH(MIN(B7:D7),B7:D7,0))'
    # clearer
    cf["A10"] = "Recommend"
    cf["B10"] = '=IF(AND(B7<=C7,B7<=D7),"Vendor A",IF(C7<=D7,"Vendor B","Vendor C"))'

    rec = wb.create_sheet("Recommendation")
    rec["A1"] = "采购建议（≤500字）"
    rec["A1"].font = Font(bold=True, size=14)
    rec["A3"] = (
        "财务口径选 Vendor A（NPV 约 1,426.8 千美元最低）。不可只看名义许可：B 许可看起来最低（300 起），"
        "但停机 80+迁移 50+支持 65 把 Y1 推到 495，NPV 反高于 A。"
        "C 许可和支持最好看、残值 120 有退出回收，但 t0 实施 300 最重，NPV 最高。"
        "非价格：A 仅两个同规模客户，C 有六个——实施风险 A 高于 C，与 NPV 排序相反。"
        "B 许可逐年上涨且停机假设最不确定，不作主选。"
        "建议：以 A 为首选进入谈判，合同写入停机 SLA 与实施验收；并行对 C 做实施风险尽调。"
        "若尽调显示 A 交付失败成本接近停机差额，再切换 C。不选 B，除非停机成本能被合同 SLA 实质压到接近 25。"
        "敏感性：若 B 停机从 80 降到 25，NPV 可低于 A——故 B 排序高度依赖最不确定假设。"
    )
    rec["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    rec.row_dimensions[3].height = 140
    rec.column_dimensions["A"].width = 110
    chart = BarChart()
    chart.type = "col"
    chart.title = "NPV at 8% ($000s)"
    chart.add_data(Reference(cf, min_col=2, min_row=7, max_col=4, max_row=7), from_rows=True, titles_from_data=False)
    chart.set_categories(Reference(cf, min_col=2, min_row=1, max_col=4, max_row=1))
    cf.add_chart(chart, "A12")
    autosize(inp, [22, 14, 14, 14, 18])
    autosize(cf, [16, 14, 14, 14])
    wb.save(OUT / "T35_Vendor_TCO_NPV.xlsx")


def t36():
    wb = Workbook()
    p = wb.active
    p.title = "Plan"
    p.append(["ID", "Task", "Duration_wd", "Depends", "Owner", "Start", "Finish"])
    style_header(p, 1, 7)
    # Dates as values so the file opens with a complete schedule; formulas documented
    from datetime import date
    rows = [
        ("A", "Requirements", 3, "none", "Product", date(2026, 9, 1), date(2026, 9, 3)),
        ("B", "Solution design", 4, "A", "Engineering", date(2026, 9, 4), date(2026, 9, 9)),
        ("C", "Legal review", 5, "A", "Legal", date(2026, 9, 4), date(2026, 9, 10)),
        ("D", "Build", 8, "B", "Engineering", date(2026, 9, 10), date(2026, 9, 21)),
        ("E", "Content", 6, "B", "Marketing", date(2026, 9, 10), date(2026, 9, 17)),
        ("F", "UAT", 4, "D,E", "Product", date(2026, 9, 22), date(2026, 9, 25)),
        ("G", "Training", 3, "C,F", "Enablement", date(2026, 9, 28), date(2026, 9, 30)),
        ("H", "Launch", 1, "G", "Program Lead", date(2026, 10, 1), date(2026, 10, 1)),
    ]
    for r in rows:
        p.append(list(r))
        p.cell(p.max_row, 6).number_format = "YYYY-MM-DD"
        p.cell(p.max_row, 7).number_format = "YYYY-MM-DD"
    p["A12"] = "Rule: successor starts next working day after ALL predecessors finish. Weekends off. No public holidays."
    p["A13"] = "Critical path: A-B-D-F-G-H = 3+4+8+4+3+1 = 23 working days. Launch 2026-10-01 HIT."
    p["A14"] = "Start_check_weekday_Sep1"
    p["B14"] = '=TEXT(F2,"DDD")'
    p["A15"] = "Launch_is_Oct1"
    p["B15"] = '=IF(F9=DATE(2026,10,1),"YES","NO")'

    g = wb.create_sheet("Gantt")
    g.append(["ID", "Task"] + [f"d{i}" for i in range(1, 24)])
    # Simplified labeled gantt with dates as header
    headers = ["ID", "Task", "2026-09-01", "09-02", "09-03", "09-04", "09-07", "09-08", "09-09",
               "09-10", "09-11", "09-14", "09-15", "09-16", "09-17", "09-18", "09-21",
               "09-22", "09-23", "09-24", "09-25", "09-28", "09-29", "09-30", "10-01"]
    g.delete_rows(1)
    g.append(headers)
    style_header(g, 1, len(headers))
    bars = {
        "A": [1, 2, 3],
        "B": [4, 5, 6, 7],
        "C": [4, 5, 6, 7, 8],
        "D": [8, 9, 10, 11, 12, 13, 14, 15],
        "E": [8, 9, 10, 11, 12, 13],
        "F": [16, 17, 18, 19],
        "G": [20, 21, 22],
        "H": [23],
    }
    names = dict(rows[i][:2] for i in range(8)) if False else {r[0]: r[1] for r in rows}
    for i, tid in enumerate(["A", "B", "C", "D", "E", "F", "G", "H"], start=2):
        g.cell(i, 1, tid)
        g.cell(i, 2, names[tid])
        for col in range(3, 26):
            day_idx = col - 2
            if day_idx in bars[tid]:
                g.cell(i, col, "█")
                g.cell(i, col).fill = PatternFill("solid", fgColor="1F6F8B")
                g.cell(i, col).font = Font(color="1F6F8B")
            else:
                g.cell(i, col, "")
    g.sheet_view.showGridLines = True

    raci = wb.create_sheet("RACI")
    raci.append(["Task", "Product", "Eng", "Legal", "Mkt", "Enablement", "Program Lead", "COO", "InfoSec"])
    style_header(raci, 1, 9)
    raci_rows = [
        ("A Requirements", "A/R", "C", "I", "I", "I", "I", "I", "I"),
        ("B Design", "C", "A/R", "I", "C", "I", "I", "I", "C"),
        ("C Legal", "I", "I", "A/R", "I", "I", "I", "C", "I"),
        ("D Build", "C", "A/R", "I", "I", "I", "I", "I", "C"),
        ("E Content", "C", "I", "C", "A/R", "I", "I", "I", "I"),
        ("F UAT", "A/R", "R", "I", "C", "C", "C", "I", "C"),
        ("G Training", "C", "I", "I", "C", "A/R", "I", "I", "I"),
        ("H Launch", "C", "C", "C", "C", "C", "A/R", "C", "I"),
    ]
    for row in raci_rows:
        raci.append(list(row))
    raci["A12"] = "Exactly one Accountable (A) per task. InfoSec Consulted during UAT. Sponsor=COO. Marketing 14-15 Sep covered by deputy; A stays with Marketing."

    risk = wb.create_sheet("Risk_Register")
    risk.append(["ID", "Risk", "Impact", "Mitigation"])
    style_header(risk, 1, 4)
    risk.append(["R1", "Legal delay 3 days", "C ends 15 Sep, still before F ends 25 Sep; not on critical path unless delay exceeds F", "Check C on 10 Sep"])
    risk.append(["R2", "Only one UAT environment", "F cannot split; 4 days no backup env", "Exclusive calendar before 21 Sep"])
    risk.append(["R3", "Marketing 14-15 Sep leave", "E in progress those days", "Deputy covers; sign-off before 13 Sep"])
    risk.append(["R4", "Zero buffer to 1 Oct", "Any +1 working day misses launch", "Daily standup on D/F; freeze scope at end of B"])
    risk.append(["R5", "Public holidays ignored per brief", "—", "Per task assumption"])
    autosize(p, [8, 22, 14, 12, 16, 14, 14])
    autosize(raci, [18] + [12] * 8)
    autosize(risk, [8, 28, 70, 40])
    for col in range(3, 26):
        g.column_dimensions[get_column_letter(col)].width = 6
    g.column_dimensions["A"].width = 6
    g.column_dimensions["B"].width = 20
    wb.save(OUT / "T36_Launch_Plan.xlsx")


def t50_xlsx():
    wb = Workbook()
    k = wb.active
    k.title = "KPI"
    k.append(["Metric", "Q2", "Q1", "Q2_LY"])
    style_header(k, 1, 4)
    k.append(["ARR_m", 120, 110, 95])
    k.append(["NRR", 0.04 + 1, 1.10, 1.12])  # wait NRR is 104% = 1.04
    # fix: 104% should be 1.04
    k["B3"] = 1.04
    k["C3"] = 1.10
    k["D3"] = 1.12
    k.append(["Gross_Margin", 0.68, 0.72, 0.70])
    k.append(["Logo_Churn", 0.028, 0.019, 0.017])
    k.append(["New_Logos", 40, 55, 48])
    k.append(["Cash_m", 28, 31, 36])
    k.append(["Monthly_Burn_m", 3.0, 2.5, 2.2])
    k["B3"].number_format = "0.00%"
    # 104% as percent: 1.04 formatted 0.00% shows 104.00%  -- actually 1.04 with 0% is 104%. Good.
    for cell in ["C3", "D3", "B4", "C4", "D4", "B5", "C5", "D5"]:
        k[cell].number_format = "0.00%"
    # NRR 104% stored as 1.04 - wait I set k.append NRR wrong then overwrote B3. Row 3 is NRR. Row 4 is GM from append after overwrite... 
    # Let me not use the broken append. I'll rebuild KPI cleanly.

    wb.remove(k)
    k = wb.create_sheet("KPI", 0)
    k.append(["Metric", "Q2", "Q1", "Q2_LY", "Unit"])
    style_header(k, 1, 5)
    metrics = [
        ("ARR_m", 120, 110, 95, "$m"),
        ("NRR", 1.04, 1.10, 1.12, "ratio"),
        ("Gross_Margin", 0.68, 0.72, 0.70, "%"),
        ("Logo_Churn", 0.028, 0.019, 0.017, "%"),
        ("New_Logos", 40, 55, 48, "count"),
        ("Cash_m", 28, 31, 36, "$m"),
        ("Monthly_Burn_m", 3.0, 2.5, 2.2, "$m"),
    ]
    for row in metrics:
        k.append(list(row))
    for r, fmt in [(3, "0.00%"), (4, "0.00%"), (5, "0.00%")]:
        for c in range(2, 5):
            k.cell(r, c).number_format = fmt
    # NRR as 104% : 1.04 with 0.00% = 104.00%. Good.

    c = wb.create_sheet("Calculations")
    c.append(["Calc", "Formula_result", "Excel"])
    style_header(c, 1, 3)
    calcs = [
        ("ARR_QoQ", "= (KPI!B2-KPI!C2)/KPI!C2", "0.00%"),
        ("ARR_YoY", "= (KPI!B2-KPI!D2)/KPI!D2", "0.00%"),
        ("NRR_pp_vs_Q1", "=KPI!B3-KPI!C3", "0.00%"),
        ("GM_pp_vs_Q1", "=KPI!B4-KPI!C4", "0.00%"),
        ("Churn_pp_vs_Q1", "=KPI!B5-KPI!C5", "0.00%"),
        ("NewLogos_QoQ", "=(KPI!B6-KPI!C6)/KPI!C6", "0.00%"),
        ("Cash_QoQ", "=(KPI!B7-KPI!C7)/KPI!C7", "0.00%"),
        ("Runway_Q2", "=KPI!B7/KPI!B8", "0.000"),
        ("Runway_Q1", "=KPI!C7/KPI!C8", "0.000"),
        ("Price_uplift_theoretical", "=0.08*KPI!B2", "0.00"),
        ("Onboarding_cost", 0.36, "0.00"),
        ("SLA_eng_cost", 0.25, "0.00"),
        ("Combined_D1D2", "=B12+B13", "0.00"),
    ]
    # Write formulas in column B
    c["A2"] = "ARR_QoQ"; c["B2"] = "=(KPI!B2-KPI!C2)/KPI!C2"; c["C2"] = "=(KPI!B2-KPI!C2)/KPI!C2"; c["B2"].number_format = "0.00%"
    c["A3"] = "ARR_YoY"; c["B3"] = "=(KPI!B2-KPI!D2)/KPI!D2"; c["B3"].number_format = "0.00%"
    c["A4"] = "NRR_pp_vs_Q1"; c["B4"] = "=KPI!B3-KPI!C3"; c["B4"].number_format = "0.00%"
    c["A5"] = "GM_pp_vs_Q1"; c["B5"] = "=KPI!B4-KPI!C4"; c["B5"].number_format = "0.00%"
    c["A6"] = "Churn_pp_vs_Q1"; c["B6"] = "=KPI!B5-KPI!C5"; c["B5"].number_format = "0.00%"
    c["A7"] = "NewLogos_QoQ"; c["B7"] = "=(KPI!B6-KPI!C6)/KPI!C6"; c["B7"].number_format = "0.00%"
    c["A8"] = "Cash_QoQ"; c["B8"] = "=(KPI!B7-KPI!C7)/KPI!C7"; c["B8"].number_format = "0.00%"
    c["A9"] = "Runway_Q2"; c["B9"] = "=KPI!B7/KPI!B8"; c["B9"].number_format = "0.000"
    c["A10"] = "Runway_Q1"; c["B10"] = "=KPI!C7/KPI!C8"; c["B10"].number_format = "0.000"
    c["A11"] = "Price_uplift_theoretical"; c["B11"] = "=0.08*KPI!B2"; c["B11"].number_format = "0.00"
    c["A12"] = "Onboarding_cost"; c["B12"] = 0.36
    c["A13"] = "SLA_eng_cost"; c["B13"] = 0.25
    c["A14"] = "Combined_D1D2"; c["B14"] = "=B12+B13"
    c["A6"].number_format = "@"
    c["B6"].number_format = "0.00%"
    c["A16"] = "All self-calculated figures in the memo/PPT are sourced from this sheet."

    d = wb.create_sheet("Decision_Scenarios")
    d.append(["ID", "Decision", "Spend", "Benefit_claimed", "Evidence", "Board_vote"])
    style_header(d, 1, 6)
    d.append(["D1", "2 onboarding managers", "0.36m/yr (=Calculations!B12)", "10→7 weeks", "unpiloted past-project estimate", "Approve as pilot"])
    d.append(["D2", "Premium SLA build", "0.25m once, earliest Q4 (=Calculations!B13)", "productize 99.95%", "pricing unknown; 2 logos willing", "Approve build; no customer commit"])
    d.append(["D3", "Raise all prices 8%", 0, "theoretical +9.6m ARR (=Calculations!B11)", "no elasticity; conflicts with churn/NRR", "Reject"])
    d["B6"] = "Linked spend checks"
    d["C6"] = "=Calculations!B12"
    d["D6"] = "=Calculations!B13"
    d["E6"] = "=Calculations!B11"
    autosize(k, [18, 12, 12, 12, 12])
    autosize(c, [28, 18, 40])
    autosize(d, [8, 28, 42, 36, 40, 36])
    wb.save(OUT / "T50_Board_KPI_Appendix.xlsx")


if __name__ == "__main__":
    t11(); print("T11")
    t12(); print("T12")
    t13(); print("T13")
    t14(); print("T14")
    t15(); print("T15")
    t34(); print("T34")
    t35(); print("T35")
    t36(); print("T36")
    t50_xlsx(); print("T50 xlsx")
