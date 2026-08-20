#!/usr/bin/env python3
"""DOCX deliverables: T22 MSA, T50 board memo."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = Path("/workspace/office-deliverables")
OUT.mkdir(parents=True, exist_ok=True)


def set_cn_font(run, name="Microsoft YaHei", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_p(doc, text, size=11, bold=False, space_after=8, font="Calibri"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def t22():
    doc = Document()
    add_p(doc, "MASTER SERVICES AGREEMENT", 16, True, 4)
    add_p(doc, "(Working draft for counsel review — not for signature)", 11, True, 12)
    add_p(doc, "This Master Services Agreement (“Agreement”) is made as of the Effective Date [TO CONFIRM] between Orion Analytics Ltd. (“Supplier”) and Northstar Retail Inc. (“Customer”). Addresses and notice emails: [TO CONFIRM].")
    add_p(doc, "This is a working draft only. It must be reviewed by qualified counsel before signing. Data protection, insurance, export control and similar topics are not agreed and are not stated as agreed.")

    sections = [
        ("1. Definitions",
         "1.1 “SOW” means a signed Statement of Work describing data analytics configuration and support.\n"
         "1.2 “Deliverables” means items specifically commissioned under an SOW.\n"
         "1.3 “Background IP” means IP owned or licensed by a party independently of this Agreement.\n"
         "1.4 “Confidential Information” has the meaning in Section 8.\n"
         "1.5 Other capitalized terms have the meanings given in the term sheet or the Section where first used."),
        ("2. Structure and precedence",
         "Services are provided under this Agreement and signed SOWs. If an SOW conflicts with this Agreement on commercial details (fees, timeline, scope) the SOW governs for that engagement; this Agreement governs all other terms. No SOW amends Sections 7–13 unless the SOW expressly identifies the Section."),
        ("3. Services",
         "Supplier shall perform data analytics configuration and support as described in signed SOWs. Supplier shall perform in a professional manner consistent with Section 9."),
        ("4. Term and renewal",
         "The term is two (2) years from the Effective Date [TO CONFIRM], then automatic successive one-year renewals unless either party gives at least sixty (60) days’ notice of non-renewal before the then-current term ends."),
        ("5. SOWs and change control",
         "5.1 Each SOW shall set fees, scope and timeline.\n"
         "5.2 Changes to scope, timeline or fees require a written Change Order signed by both parties. Verbal instructions do not amend an SOW."),
        ("6. Fees and payment",
         "6.1 Fees are as stated in each SOW.\n"
         "6.2 Invoices are issued monthly in arrears. Amounts are due thirty (30) days after invoice date.\n"
         "6.3 Disputed amounts must be raised within fifteen (15) days of invoice date. Timely disputed amounts may be withheld pending resolution; undisputed amounts remain payable.\n"
         "6.4 Taxes, late-interest rate and invoice currency: [TO CONFIRM]."),
        ("7. Intellectual property",
         "7.1 Each party retains its Background IP.\n"
         "7.2 Upon full payment of the applicable SOW fees, Customer owns specifically commissioned Deliverables.\n"
         "7.3 Supplier retains general tools, templates and know-how. To the extent Deliverables embed Supplier Background IP, Supplier grants Customer a perpetual, non-exclusive, internal-use license to that embedded Background IP.\n"
         "7.4 Except as stated, neither party assigns residual know-how that is not a Deliverable."),
        ("8. Confidentiality",
         "8.1 Mutual. Each party shall protect the other party’s Confidential Information for three (3) years after disclosure; trade secrets shall be protected while legally protected as trade secrets.\n"
         "8.2 Use is limited to performing this Agreement.\n"
         "8.3 Standard exclusions (public, previously known, independently developed, lawfully received from a third party) apply. Compelled disclosure: [TO CONFIRM process].\n"
         "8.4 Definition of Confidential Information: [TO CONFIRM if marking required]."),
        ("9. Warranty",
         "9.1 Supplier warrants that services will be performed professionally.\n"
         "9.2 Customer’s exclusive remedy for breach of this Section 9 is re-performance, provided Customer notifies Supplier within thirty (30) days after the non-conforming services were performed.\n"
         "9.3 No other warranty is agreed in the term sheet. Fitness, merchantability and error-free operation are not stated as agreed."),
        ("10. Indemnity",
         "10.1 Supplier shall defend Customer against third-party claims that Deliverables infringe IP, and pay finally awarded damages, subject to Section 11.\n"
         "10.2 Exclusions: Customer modifications; combinations not supplied by Supplier; use outside documentation.\n"
         "10.3 Customer shall give prompt notice and reasonable cooperation. Defense-control mechanics are [TO CONFIRM] because they were not in the term sheet.\n"
         "10.4 No Customer indemnity was agreed."),
        ("11. Liability",
         "11.1 Each party’s aggregate liability under this Agreement is capped at the fees paid or payable in the twelve (12) months preceding the claim.\n"
         "11.2 The cap does not apply to: fraud; wilful misconduct; confidentiality breach; or IP infringement indemnity under Section 10.\n"
         "11.3 Neither party is liable for indirect or consequential loss, except amounts payable under an indemnity.\n"
         "11.4 Nothing in this Section excludes liability that cannot be excluded under applicable law [TO CONFIRM counsel]."),
        ("12. Termination",
         "12.1 Either party may terminate this Agreement or an affected SOW for material breach if the breach remains uncured thirty (30) days after written notice.\n"
         "12.2 Either party may terminate immediately for insolvency.\n"
         "12.3 Customer may terminate an SOW for convenience with thirty (30) days’ notice, and shall pay committed fees plus non-cancellable costs.\n"
         "12.4 Survival: Sections 7, 8, 10, 11, 13 and accrued payment obligations survive."),
        ("13. Governing law and disputes",
         "13.1 Singapore law.\n"
         "13.2 Disputes shall be escalated to the parties’ executives for fifteen (15) days.\n"
         "13.3 Thereafter, disputes shall be finally resolved by SIAC arbitration in Singapore, English language, one arbitrator."),
        ("14. Notices",
         "Notices under this Agreement shall be in writing to the emails/addresses [TO CONFIRM]."),
        ("15. General",
         "15.1 Assignment is [TO CONFIRM] (term sheet silent).\n"
         "15.2 This Agreement and signed SOWs are the entire agreement on the subject matter.\n"
         "15.3 Amendments must be in writing and signed.\n"
         "15.4 Counterparts and electronic signatures permitted [TO CONFIRM]."),
        ("16. Records and audit of fees",
         "Supplier shall keep reasonable records supporting invoices under each SOW for twelve (12) months after the SOW ends. Customer may inspect those fee records on reasonable notice, not more than once per calendar year. This Section does not create a security audit right."),
        ("17–22. Not agreed",
         "Non-solicitation, force majeure, publicity, data protection addendum, insurance limits and export controls are NOT agreed and are not granted by this draft. If Customer requires processing of personal data, the parties shall negotiate a separate addendum before such processing starts."),
        ("Items to confirm",
         "Effective Date; addresses; notice emails; DPA; insurance; invoice currency/tax/late interest; confidentiality marking; assignment; indemnity defense control; electronic signature protocol."),
    ]
    for title, body in sections:
        add_p(doc, title, 13, True, 6)
        add_p(doc, body, 11, False, 10)
    add_p(doc, "Signature blocks", 13, True)
    add_p(doc, "For Orion Analytics Ltd.                    For Northstar Retail Inc.")
    add_p(doc, "Name / Title / Date: [TO CONFIRM]           Name / Title / Date: [TO CONFIRM]")
    add_p(doc, "Counsel legend: This working draft implements the signed term sheet. It must be reviewed by qualified counsel before signature. Do not treat bracketed items as agreed.", 10, True)
    doc.save(OUT / "T22_Orion_Northstar_MSA_Working_Draft.docx")


def t50_memo():
    doc = Document()
    p = add_p(doc, "致董事会｜经营与三项决定（不超过两页）", 16, True, 8)
    add_p(doc, "事实 / 管理层判断 / 建议分开。不承诺 AI 模块日期、不承诺 SLA、不承诺提价。三件材料数字与 XLSX 一致。", 10, False, 10)

    add_p(doc, "三项决定", 13, True, 6)
    add_p(doc, "D1 批准增加 2 名上线经理（年成本 0.36m）作为试点编制，不把“10 周→7 周”当已验证承诺。")
    add_p(doc, "D2 批准高级 SLA 产品化工程 0.25m，Q4 最早可上线；定价待定，不向客户承诺 99.95% 直至产品上线且合同条款完成。")
    add_p(doc, "D3 否决全体客户立即提价 8%。允许仅对自愿付费的高级 SLA 客户单独定价（价格未定）。")

    add_p(doc, "事实（经营）", 13, True, 6)
    add_p(doc, "ARR 120 vs Q1 110 vs LY 95。NRR 104% vs 110% / 112%。毛利率 68% vs 72% / 70%。Logo churn 2.8% vs 1.9% / 1.7%。New logos 40 vs 55 / 48。现金 28 vs 31 / 36。月烧 3.0 vs 2.5 / 2.2。")
    add_p(doc, "测算：ARR 环比 +9.1%、同比 +26.3%；NRR −6 个百分点（对 Q1）；毛利率 −4 个百分点；流失 +0.9 个百分点；新客 logo −27.3%；现金 −3m（−9.7%）；烧钱 +0.5。简单跑道 Cash/Burn：Q2 9.3 个月 vs Q1 12.4 个月。")

    add_p(doc, "事实（客户/产品/法务）", 13, True, 6)
    add_p(doc, "五家企业要 99.95% SLA，当前标准 99.9%；其中两家愿为高级 SLA 付费。近 10 个流失访谈：6 上线慢、3 AI 摘要准确性、1 价格。销售建议全体提价 8%，称可立即增加 ARR，未提供流失弹性。理论年化毛增量 0.08×120=9.6m，未考虑流失、折扣、时点。")
    add_p(doc, "AI 模块内部目标 9 月 15 日；销售材料曾写 8 月 15 日，产品未批准。离线测试准确率 88%，上线门槛 92%。法务：自动续约通知仅提前 30 天，两地区内部政策要求至少 60 天，须在下轮续约前改流程。")

    add_p(doc, "管理层判断（非因果证明）", 13, True, 6)
    add_p(doc, "流失访谈样本 10，上线慢是最频繁提及项，可能与 NRR/流失恶化同期，但非统计证明。AI 准确性被 3 家提及，与 88%<92% 一致，支持暂缓对外日期。提价 8% 在流失已升、NRR 已降时风险高于销售叙事。")

    add_p(doc, "建议与表决", 13, True, 6)
    add_p(doc, "见文首 D1–D3。现金 28、烧 3.0，D1+D2 年化/一次性合计 0.61m，不改变跑道数量级，但 D3 的 9.6m 为未风险调整理论值，不能当预算。请表决：D1 是/否；D2 是/否（含“不得对客户承诺日期与 SLA”附带条件）；D3 否决立即全量提价 是/否。另：授权法务在下轮续约前把通知改到 ≥60 天（合规项，非增长实验）。")
    add_p(doc, "附件：T50_Board_Pack.pptx（6 页）；T50_Board_KPI_Appendix.xlsx（KPI / Calculations / Decision_Scenarios）。", 10, False, 6)
    doc.save(OUT / "T50_Board_Memo.docx")


if __name__ == "__main__":
    t22()
    print("T22")
    t50_memo()
    print("T50 memo")
