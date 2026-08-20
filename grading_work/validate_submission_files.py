#!/usr/bin/env python3
"""Independently validate candidate submission files without modifying them."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
SUBMISSIONS = ROOT / "submission_bundle"
RESULTS = ROOT / "grading_results"


def validate_zip_container(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        corrupt_member = archive.testzip()
        member_count = len(archive.namelist())
    return {
        "zip_crc_valid": corrupt_member is None,
        "corrupt_member": corrupt_member,
        "container_members": member_count,
    }


def validate_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    formulas = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas.append(f"{sheet.title}!{cell.coordinate}")
    return {
        "type": "xlsx",
        "sheets": workbook.sheetnames,
        "sheet_count": len(workbook.sheetnames),
        "formula_count": len(formulas),
        "formula_cells_sample": formulas[:20],
        "cached_sheet_count": len(cached.sheetnames),
    }


def validate_pptx(path: Path) -> dict:
    presentation = Presentation(path)
    slide_text = []
    chart_count = 0
    table_count = 0
    for index, slide in enumerate(presentation.slides, start=1):
        text = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text.extend(
                    paragraph.text
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
            if getattr(shape, "has_chart", False):
                chart_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
        slide_text.append({"slide": index, "text": " | ".join(text)[:500]})
    return {
        "type": "pptx",
        "slide_count": len(presentation.slides),
        "chart_count": chart_count,
        "table_count": table_count,
        "slide_text": slide_text,
    }


def validate_docx(path: Path) -> dict:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    return {
        "type": "docx",
        "paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "text_sample": "\n".join(paragraphs)[:2000],
    }


def main() -> None:
    RESULTS.mkdir(exist_ok=True)
    records = []
    for path in sorted(SUBMISSIONS.rglob("*")):
        if not path.is_file():
            continue
        record = {
            "path": str(path.relative_to(ROOT)),
            "size_bytes": path.stat().st_size,
            "valid": True,
            "error": "",
        }
        try:
            if path.suffix.lower() in {".xlsx", ".pptx", ".docx"}:
                record.update(validate_zip_container(path))
            if path.suffix.lower() == ".xlsx":
                record.update(validate_xlsx(path))
            elif path.suffix.lower() == ".pptx":
                record.update(validate_pptx(path))
            elif path.suffix.lower() == ".docx":
                record.update(validate_docx(path))
            elif path.suffix.lower() in {".md", ".csv"}:
                text = path.read_text(encoding="utf-8")
                record.update(
                    {
                        "type": path.suffix.lower().lstrip("."),
                        "character_count": len(text),
                        "line_count": len(text.splitlines()),
                    }
                )
            else:
                record["type"] = "other"
        except Exception as exc:
            record["valid"] = False
            record["error"] = f"{type(exc).__name__}: {exc}"
        records.append(record)

    payload = {
        "files_checked": len(records),
        "valid_files": sum(record["valid"] for record in records),
        "invalid_files": sum(not record["valid"] for record in records),
        "records": records,
    }
    output = RESULTS / "submission_file_validation.json"
    output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {key: payload[key] for key in ("files_checked", "valid_files", "invalid_files")},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
