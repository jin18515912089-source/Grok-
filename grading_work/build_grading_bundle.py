#!/usr/bin/env python3
"""Validate domain grading JSON and build the Batch-50 grading bundle."""

from __future__ import annotations

import csv
import json
import math
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT.parent / "grading_results"
PART_FILES = [
    ROOT / "T01-T05.json",
    ROOT / "T06-T10.json",
    ROOT / "T11-T15.json",
    ROOT / "T16-T20.json",
    ROOT / "T21-T25.json",
    ROOT / "T26-T30.json",
    ROOT / "T31-T35.json",
    ROOT / "T36-T40.json",
    ROOT / "T41-T45.json",
    ROOT / "T46-T50.json",
]

DOMAIN_BY_RANGE = {
    range(1, 6): "管理沟通与写作",
    range(6, 11): "演示与可视化",
    range(11, 16): "电子表格与数据处理",
    range(16, 21): "营销、销售与客户",
    range(21, 26): "法律、合规与政策",
    range(26, 31): "研究、分析与尽调",
    range(31, 36): "财务、投资与会计",
    range(36, 41): "运营、项目与采购",
    range(41, 46): "人力资源与内部沟通",
    range(46, 51): "综合代理式白领任务",
}
DOMAINS = list(DOMAIN_BY_RANGE.values())
TASK_FIELDS = [
    "task_id",
    "domain",
    "submission_status",
    "required_files",
    "actual_files",
    "file_validity",
    "complete_deliverables",
    "subtotal",
    "cap_triggered",
    "cap_rule",
    "final_score",
    "top_issues",
    "strengths",
    "confidence",
    "human_review_flag",
    "human_review_reason",
]


def expected_domain(task_number: int) -> str:
    for task_range, domain in DOMAIN_BY_RANGE.items():
        if task_number in task_range:
            return domain
    raise ValueError(f"Unexpected task number {task_number}")


def close(a: float, b: float) -> bool:
    return math.isclose(a, b, abs_tol=0.051)


def load_and_validate() -> tuple[list[dict], list[dict]]:
    errors: list[str] = []
    warnings: list[str] = []
    tasks: list[dict] = []

    missing_parts = [str(path.name) for path in PART_FILES if not path.exists()]
    if missing_parts:
        raise FileNotFoundError(f"Missing grading parts: {', '.join(missing_parts)}")

    for path in PART_FILES:
        with path.open(encoding="utf-8") as handle:
            payload = json.load(handle)
        if set(payload) != {"tasks"} or not isinstance(payload["tasks"], list):
            errors.append(f"{path.name}: top-level object must contain only tasks[]")
            continue
        tasks.extend(payload["tasks"])

    ids = [task.get("task_id") for task in tasks]
    expected_ids = [f"T{i:02d}" for i in range(1, 51)]
    duplicates = sorted(task_id for task_id, count in Counter(ids).items() if count > 1)
    missing = sorted(set(expected_ids) - set(ids))
    extra = sorted(set(ids) - set(expected_ids))
    if len(tasks) != 50:
        errors.append(f"Expected 50 task rows, found {len(tasks)}")
    if duplicates:
        errors.append(f"Duplicate task IDs: {duplicates}")
    if missing:
        errors.append(f"Missing task IDs: {missing}")
    if extra:
        errors.append(f"Unexpected task IDs: {extra}")

    for task in tasks:
        task_id = task.get("task_id", "<missing>")
        for field in TASK_FIELDS:
            if field not in task:
                errors.append(f"{task_id}: missing field {field}")
        match = re.fullmatch(r"T(\d{2})", str(task_id))
        if not match:
            continue
        number = int(match.group(1))
        if task.get("domain") != expected_domain(number):
            errors.append(
                f"{task_id}: domain {task.get('domain')!r} != {expected_domain(number)!r}"
            )
        if task.get("submission_status") not in {"COMPLETE", "PARTIAL", "FAILED"}:
            errors.append(f"{task_id}: invalid submission_status")
        if task.get("file_validity") not in {"VALID", "PARTIAL", "INVALID"}:
            errors.append(f"{task_id}: invalid file_validity")
        if task.get("confidence") not in {"HIGH", "MEDIUM", "LOW"}:
            errors.append(f"{task_id}: invalid confidence")
        if task.get("human_review_flag") not in {"YES", "NO"}:
            errors.append(f"{task_id}: invalid human_review_flag")
        if not isinstance(task.get("complete_deliverables"), bool):
            errors.append(f"{task_id}: complete_deliverables must be boolean")
        if not isinstance(task.get("cap_triggered"), bool):
            errors.append(f"{task_id}: cap_triggered must be boolean")

        dimensions = task.get("dimensions")
        if not isinstance(dimensions, list) or not dimensions:
            errors.append(f"{task_id}: dimensions must be a nonempty list")
            continue
        dimension_max = 0.0
        dimension_score = 0.0
        dimension_names: set[str] = set()
        for index, dimension in enumerate(dimensions, start=1):
            required = {
                "dimension",
                "max_points",
                "score",
                "evidence",
                "deduction_reason",
            }
            absent = required - set(dimension)
            if absent:
                errors.append(f"{task_id} dimension {index}: missing {sorted(absent)}")
                continue
            name = str(dimension["dimension"]).strip()
            if not name:
                errors.append(f"{task_id} dimension {index}: empty name")
            if name in dimension_names:
                errors.append(f"{task_id}: duplicate dimension {name!r}")
            dimension_names.add(name)
            maximum = float(dimension["max_points"])
            score = float(dimension["score"])
            dimension_max += maximum
            dimension_score += score
            if maximum <= 0 or score < 0 or score > maximum + 0.051:
                errors.append(
                    f"{task_id} {name}: invalid score {score}/{maximum}"
                )
            if not close(score * 10, round(score * 10)):
                errors.append(f"{task_id} {name}: score is not in 0.1 increments")
            if not str(dimension["evidence"]).strip():
                errors.append(f"{task_id} {name}: empty evidence")
            if not str(dimension["deduction_reason"]).strip():
                warnings.append(f"{task_id} {name}: empty deduction reason")

        subtotal = float(task.get("subtotal", -1))
        final_score = float(task.get("final_score", -1))
        if not close(dimension_max, 10.0):
            errors.append(f"{task_id}: dimension max sum {dimension_max:.2f} != 10")
        if not close(subtotal, round(dimension_score, 1)):
            errors.append(
                f"{task_id}: subtotal {subtotal:.1f} != dimension sum {dimension_score:.1f}"
            )
        if not 0 <= final_score <= subtotal + 0.051 <= 10.051:
            errors.append(
                f"{task_id}: expected 0 <= final_score <= subtotal <= 10, "
                f"found {final_score}/{subtotal}"
            )
        if len(task.get("top_issues", [])) != 3:
            warnings.append(f"{task_id}: top_issues count is not 3")
        if len(task.get("strengths", [])) != 2:
            warnings.append(f"{task_id}: strengths count is not 2")
        if task.get("cap_triggered") and not str(task.get("cap_rule", "")).strip():
            errors.append(f"{task_id}: cap triggered without cap rule")

    domain_counts = Counter(task.get("domain") for task in tasks)
    for domain in DOMAINS:
        if domain_counts[domain] != 5:
            errors.append(f"{domain}: expected 5 tasks, found {domain_counts[domain]}")

    if errors:
        raise ValueError("Validation failed:\n- " + "\n- ".join(errors))
    tasks.sort(key=lambda task: task["task_id"])
    return tasks, warnings


def score_stats(tasks: list[dict]) -> dict:
    domain_scores = {}
    for domain in DOMAINS:
        domain_tasks = [task for task in tasks if task["domain"] == domain]
        domain_scores[domain] = round(
            sum(float(task["final_score"]) for task in domain_tasks) / 5 * 10, 1
        )
    complete = sum(bool(task["complete_deliverables"]) for task in tasks)
    caps = sum(bool(task["cap_triggered"]) for task in tasks)
    reviews = sum(task["human_review_flag"] == "YES" for task in tasks)
    return {
        "overall": round(sum(float(task["final_score"]) for task in tasks) / 50 * 10, 1),
        "completion": complete / 50,
        "complete_count": complete,
        "severe_error_rate": caps / 50,
        "cap_count": caps,
        "review_count": reviews,
        "domain_scores": domain_scores,
    }


def join_list(value: list[str]) -> str:
    return "；".join(str(item) for item in value)


def style_sheet(ws, freeze: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if freeze:
        ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells
        )
        ws.column_dimensions[letter].width = min(max(max_length + 2, 10), 55)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(tasks: list[dict], stats: dict, warnings: list[str]) -> Path:
    path = OUTPUT_DIR / "GPT5.6_batch_scores.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["指标", "结果", "计算口径"])
    summary_rows = [
        ("模型匿名ID", "ANON-BATCH-001", "盲评匿名标识"),
        ("测评模式", "Batch-50", "T01–T50 一次批量评分"),
        ("已评分题数", 50, "TaskScores 非空且唯一任务数"),
        ("完成率", stats["completion"], "全部要求文件存在且可打开的任务数/50"),
        ("严重错误率/CAP率", stats["severe_error_rate"], "触发任一CAP的任务数/50"),
        ("总质量分", stats["overall"], "50题最终分平均值×10"),
        ("需人工复核题数", stats["review_count"], "human_review_flag=YES"),
    ]
    for row in summary_rows:
        summary.append(row)
    summary.append([])
    summary.append(["能力域", "领域分(0-100)", "计算口径"])
    for domain in DOMAINS:
        summary.append([domain, stats["domain_scores"][domain], "该领域5题平均分×10"])
    summary["B5"].number_format = "0.0%"
    summary["B6"].number_format = "0.0%"
    style_sheet(summary)
    summary.auto_filter.ref = f"A9:C{summary.max_row}"
    summary.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 48

    task_scores = wb.create_sheet("TaskScores")
    task_headers = [
        "task_id",
        "domain",
        "submission_status",
        "required_files",
        "actual_files",
        "file_validity",
        "complete_deliverables",
        "subtotal",
        "cap_triggered",
        "cap_rule",
        "final_score",
        "top_issues",
        "strengths",
        "confidence",
        "human_review_flag",
        "human_review_reason",
    ]
    task_scores.append(task_headers)
    for task in tasks:
        task_scores.append(
            [
                task["task_id"],
                task["domain"],
                task["submission_status"],
                task["required_files"],
                task["actual_files"],
                task["file_validity"],
                "YES" if task["complete_deliverables"] else "NO",
                float(task["subtotal"]),
                "YES" if task["cap_triggered"] else "NO",
                task["cap_rule"],
                float(task["final_score"]),
                join_list(task["top_issues"]),
                join_list(task["strengths"]),
                task["confidence"],
                task["human_review_flag"],
                task["human_review_reason"],
            ]
        )
    style_sheet(task_scores, "A2")
    task_scores.conditional_formatting.add(
        f"K2:K{task_scores.max_row}",
        CellIsRule(
            operator="lessThan",
            formula=["6"],
            fill=PatternFill("solid", fgColor="F4CCCC"),
        ),
    )

    dimensions = wb.create_sheet("DimensionScores")
    dimensions.append(
        [
            "task_id",
            "domain",
            "dimension",
            "max_points",
            "score",
            "evidence",
            "deduction_reason",
        ]
    )
    for task in tasks:
        for dimension in task["dimensions"]:
            dimensions.append(
                [
                    task["task_id"],
                    task["domain"],
                    dimension["dimension"],
                    float(dimension["max_points"]),
                    float(dimension["score"]),
                    dimension["evidence"],
                    dimension["deduction_reason"],
                ]
            )
    style_sheet(dimensions, "A2")

    qa = wb.create_sheet("QA")
    qa.append(["scope", "check", "status", "detail", "human_review_flag"])
    qa_checks = [
        ("BATCH", "TaskScores row count", "PASS", "Exactly 50 rows", "NO"),
        ("BATCH", "Task ID uniqueness", "PASS", "T01–T50, no gaps/duplicates", "NO"),
        (
            "BATCH",
            "Score bounds",
            "PASS",
            "Every task satisfies 0≤final_score≤subtotal≤10",
            "NO",
        ),
        (
            "BATCH",
            "Dimension maxima",
            "PASS",
            "Every task dimension max_points sum equals 10",
            "NO",
        ),
        (
            "BATCH",
            "Domain coverage",
            "PASS",
            "10 domains × 5 tasks",
            "NO",
        ),
        (
            "BATCH",
            "Aggregation formulas",
            "PASS",
            "Overall=mean(final_score)×10; domain=5-task mean×10",
            "NO",
        ),
    ]
    for warning in warnings:
        qa_checks.append(("BATCH", "Non-blocking validation warning", "WARN", warning, "YES"))
    for row in qa_checks:
        qa.append(row)
    for task in tasks:
        reasons = []
        if not task["complete_deliverables"]:
            reasons.append("交付文件不完整")
        if task["file_validity"] != "VALID":
            reasons.append(f"文件有效性={task['file_validity']}")
        if task["cap_triggered"]:
            reasons.append(f"CAP: {task['cap_rule']}")
        if task["confidence"] != "HIGH":
            reasons.append(f"置信度={task['confidence']}")
        if task["human_review_flag"] == "YES":
            reasons.append(task["human_review_reason"])
        qa.append(
            [
                task["task_id"],
                "Task validation/review",
                "REVIEW" if reasons else "PASS",
                "；".join(reasons) if reasons else "文件、计分和聚合校验通过",
                task["human_review_flag"],
            ]
        )
    style_sheet(qa, "A2")
    qa.conditional_formatting.add(
        f"C2:C{qa.max_row}",
        CellIsRule(
            operator="equal",
            formula=['"REVIEW"'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)
    return path


def write_scores_csv(tasks: list[dict]) -> Path:
    path = OUTPUT_DIR / "scores.csv"
    fieldnames = [
        "task_id",
        "domain",
        "submission_status",
        "file_validity",
        "complete_deliverables",
        "subtotal",
        "cap_triggered",
        "final_score",
        "top_issues",
        "strengths",
        "confidence",
        "human_review_flag",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for task in tasks:
            writer.writerow(
                {
                    "task_id": task["task_id"],
                    "domain": task["domain"],
                    "submission_status": task["submission_status"],
                    "file_validity": task["file_validity"],
                    "complete_deliverables": task["complete_deliverables"],
                    "subtotal": f"{float(task['subtotal']):.1f}",
                    "cap_triggered": task["cap_triggered"],
                    "final_score": f"{float(task['final_score']):.1f}",
                    "top_issues": join_list(task["top_issues"]),
                    "strengths": join_list(task["strengths"]),
                    "confidence": task["confidence"],
                    "human_review_flag": task["human_review_flag"],
                }
            )
    return path


def write_progress(tasks: list[dict]) -> Path:
    path = OUTPUT_DIR / "grading_progress.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["task_id", "status"])
        writer.writerows((task["task_id"], "GRADED") for task in tasks)
    return path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    document.add_paragraph()


def systemic_issue_summary(tasks: list[dict]) -> list[str]:
    patterns = {
        "缺少量化闭环或阈值": ["阈值", "量化", "验收", "责任人", "时点", "闭环"],
        "计算/关键数值偏差": ["计算", "数字", "数值", "公式", "总分", "差额"],
        "文件结构或格式不完全符合": ["页", "工作表", "格式", "文件", "字数", "图表"],
        "分析深度或论证不足": ["分析", "原因", "论证", "洞察", "因果", "证据"],
        "合规、风险或不确定性处理不足": ["合规", "风险", "隐私", "法律", "假设", "不确定"],
    }
    issue_text = "\n".join(join_list(task["top_issues"]) for task in tasks)
    ranked = []
    for label, keywords in patterns.items():
        count = sum(issue_text.count(keyword) for keyword in keywords)
        if count:
            ranked.append((count, label))
    ranked.sort(reverse=True)
    return [label for _, label in ranked[:4]]


def write_report(tasks: list[dict], stats: dict) -> Path:
    path = OUTPUT_DIR / "GPT5.6_grading_report.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor(31, 78, 121)

    title = document.add_heading("GPT-5.6 Batch-50 全量评分报告", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("匿名候选模型 ANON-BATCH-001｜测评模式 Batch-50")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("执行摘要", level=1)
    document.add_paragraph(
        f"已按评分手册与批量控制器完成 T01–T50 共 50 题逐维度评分。"
        f"总质量分为 {stats['overall']:.1f}/100；完成率为 "
        f"{stats['completion']:.1%}（{stats['complete_count']}/50）；"
        f"严重错误率为 {stats['severe_error_rate']:.1%}"
        f"（{stats['cap_count']}/50 触发 CAP）；"
        f"{stats['review_count']} 题标记为需人工复核。"
    )

    document.add_heading("领域表现", level=1)
    domain_rows = [
        [domain, f"{score:.1f}"] for domain, score in stats["domain_scores"].items()
    ]
    add_table(document, ["能力域", "领域分（0–100）"], domain_rows)
    sorted_domains = sorted(
        stats["domain_scores"].items(), key=lambda item: item[1], reverse=True
    )
    document.add_paragraph(
        f"最强领域：{sorted_domains[0][0]}（{sorted_domains[0][1]:.1f}）；"
        f"最弱领域：{sorted_domains[-1][0]}（{sorted_domains[-1][1]:.1f}）。"
    )

    document.add_heading("跨题系统性问题", level=1)
    issues = systemic_issue_summary(tasks)
    if not issues:
        issues = ["未识别出稳定的跨题问题模式，仍应按 QA 清单抽查。"]
    for issue in issues:
        document.add_paragraph(issue, style="List Bullet")

    document.add_heading("低分、CAP 与异常任务", level=1)
    exceptional = sorted(
        [
            task
            for task in tasks
            if task["final_score"] < 7
            or task["cap_triggered"]
            or task["file_validity"] != "VALID"
        ],
        key=lambda task: (task["final_score"], task["task_id"]),
    )
    if exceptional:
        add_table(
            document,
            ["任务", "最终分", "CAP", "文件有效性", "首要问题"],
            [
                [
                    task["task_id"],
                    f"{float(task['final_score']):.1f}",
                    "YES" if task["cap_triggered"] else "NO",
                    task["file_validity"],
                    task["top_issues"][0],
                ]
                for task in exceptional
            ],
        )
    else:
        document.add_paragraph("无低于 7 分、CAP 或文件有效性异常任务。")

    document.add_heading("人工复核建议", level=1)
    review_tasks = [task for task in tasks if task["human_review_flag"] == "YES"]
    if review_tasks:
        add_table(
            document,
            ["任务", "领域", "置信度", "复核原因"],
            [
                [
                    task["task_id"],
                    task["domain"],
                    task["confidence"],
                    task["human_review_reason"],
                ]
                for task in review_tasks
            ],
        )
    else:
        document.add_paragraph("没有任务被自动标记为人工复核。")
    document.add_paragraph(
        "复核优先级：先检查损坏/不可解析文件与 CAP 任务，再检查法律、财务、"
        "隐私高风险题和低置信度题。GPT 初评不替代高风险题人工判断。"
    )

    document.add_heading("逐题得分", level=1)
    add_table(
        document,
        ["任务", "领域", "小计", "最终分", "CAP", "人工复核"],
        [
            [
                task["task_id"],
                task["domain"],
                f"{float(task['subtotal']):.1f}",
                f"{float(task['final_score']):.1f}",
                "YES" if task["cap_triggered"] else "NO",
                task["human_review_flag"],
            ]
            for task in tasks
        ],
    )

    document.add_paragraph(
        "详细维度证据、扣分理由与 QA 校验见 GPT5.6_batch_scores.xlsx。"
    )
    document.save(path)
    return path


def validate_outputs(tasks: list[dict], output_paths: list[Path]) -> list[str]:
    checks: list[str] = []
    for path in output_paths:
        if not path.exists() or path.stat().st_size == 0:
            raise ValueError(f"Output missing or empty: {path}")
        checks.append(f"{path.name}: present ({path.stat().st_size} bytes)")

    csv_path = OUTPUT_DIR / "scores.csv"
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != 50 or [row["task_id"] for row in rows] != [
        f"T{i:02d}" for i in range(1, 51)
    ]:
        raise ValueError("scores.csv does not contain exactly ordered T01–T50")
    checks.append("scores.csv: exactly 50 ordered task rows")

    progress_path = OUTPUT_DIR / "grading_progress.csv"
    with progress_path.open(encoding="utf-8-sig", newline="") as handle:
        progress_rows = list(csv.DictReader(handle))
    if len(progress_rows) != 50 or any(
        row["status"] != "GRADED" for row in progress_rows
    ):
        raise ValueError("grading_progress.csv must contain 50 GRADED rows")
    checks.append("grading_progress.csv: 50 GRADED rows")

    from openpyxl import load_workbook

    workbook = load_workbook(OUTPUT_DIR / "GPT5.6_batch_scores.xlsx", read_only=True)
    if workbook.sheetnames != ["Summary", "TaskScores", "DimensionScores", "QA"]:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    if workbook["TaskScores"].max_row != 51:
        raise ValueError("TaskScores must have header plus 50 rows")
    checks.append("workbook: required sheets and 50 TaskScores rows")

    Document(OUTPUT_DIR / "GPT5.6_grading_report.docx")
    checks.append("report DOCX: opens successfully")

    for task in tasks:
        maximum = sum(float(item["max_points"]) for item in task["dimensions"])
        if not close(maximum, 10):
            raise ValueError(f"{task['task_id']} dimensions do not sum to 10")
    checks.append("dimension maxima: 10.0 for every task")
    return checks


def write_zip(output_paths: list[Path]) -> Path:
    path = OUTPUT_DIR / "GPT5.6_grading_bundle.zip"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for output in output_paths:
            archive.write(output, arcname=output.name)
    with zipfile.ZipFile(path) as archive:
        expected = {output.name for output in output_paths}
        if set(archive.namelist()) != expected or archive.testzip() is not None:
            raise ValueError("ZIP content or CRC validation failed")
    return path


def main() -> int:
    OUTPUT_DIR.mkdir(exist_ok=True)
    tasks, warnings = load_and_validate()
    stats = score_stats(tasks)
    workbook = write_workbook(tasks, stats, warnings)
    scores = write_scores_csv(tasks)
    report = write_report(tasks, stats)
    progress = write_progress(tasks)
    components = [workbook, scores, report, progress]
    checks = validate_outputs(tasks, components)
    bundle = write_zip(components)

    summary = {
        "graded_tasks": len(tasks),
        "overall_quality_score": stats["overall"],
        "completion_rate": round(stats["completion"], 4),
        "severe_error_rate": round(stats["severe_error_rate"], 4),
        "human_review_count": stats["review_count"],
        "domain_scores": stats["domain_scores"],
        "warnings": warnings,
        "checks": checks + [f"{bundle.name}: valid ZIP with four required files"],
        "bundle": str(bundle),
    }
    summary_path = OUTPUT_DIR / "validation_summary.json"
    with summary_path.open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
